"""Tests for offline resume file decoding (resume_parse + /api/parse/resume)."""

from __future__ import annotations

import base64
import io
import sys

sys.path.insert(0, "src")

from fastapi.testclient import TestClient

from agent.main import app
from agent.resume_parse import decode_resume


def _docx_bytes(text: str) -> bytes:
    import docx

    d = docx.Document()
    for ln in text.splitlines():
        d.add_paragraph(ln)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["姓名", "岗位"])
    ws.append(["张三", "后端开发"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_decode_txt():
    out = decode_resume("张三 后端工程师".encode("utf-8"), "r.txt")
    assert out["format"] == "txt" and "张三" in out["text"] and not out["error"]


def test_decode_md():
    out = decode_resume("# 李四\n- 后端".encode("utf-8"), "r.md")
    assert out["format"] == "md" and "李四" in out["text"]


def test_decode_docx():
    out = decode_resume(_docx_bytes("王五 后端工程师\n技能：golang"), "r.docx")
    assert out["format"] == "docx" and "王五" in out["text"] and "golang" in out["text"]


def test_decode_xlsx():
    out = decode_resume(_xlsx_bytes(), "r.xlsx")
    assert out["format"] == "xlsx" and "张三" in out["text"]


def test_unrecognized_binary_degrades_without_error_str():
    out = decode_resume(b"\x00\x01\x02\xff\xfe" * 20, "blob.bin")
    # Unknown extension is best-effort; must never raise and must return text field.
    assert "text" in out


def test_parse_resume_endpoint_docx():
    c = TestClient(app)
    b64 = base64.b64encode(_docx_bytes("赵六 后端工程师")).decode()
    r = c.post("/api/parse/resume", json={"data": b64, "filename": "resume.docx"})
    assert r.status_code == 200
    body = r.json()
    assert body["format"] == "docx" and "赵六" in body["text"]


def test_parse_resume_endpoint_rejects_missing():
    c = TestClient(app)
    r = c.post("/api/parse/resume", json={"filename": "x"})
    assert r.status_code == 400
